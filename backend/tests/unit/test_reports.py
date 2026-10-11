import pytest
import os
import uuid
import json
import zipfile
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi.testclient import TestClient

from app.db.session import Base, get_db
from app.main import app
from app.db.models.core import Project
from app.db.models.document_memory import Document, DocumentSheet, TitleBlockExtraction
from app.db.models.decision_memory import RuleDefinition, RuleFinding, FindingResolution
from app.db.models.operations import DecisionTrace
from app.services.reporting.service import ReportingService
from app.services.rules.engine import RuleRegistry
from app.core.config import settings

@pytest.fixture(autouse=True)
def seed_rules(db_session):
    RuleRegistry.seed_database_definitions(db_session)

def test_generate_sheet_audit_report(db_session):
    proj = Project(id=str(uuid.uuid4()), name="Torre Costanera", code="TC-01")
    db_session.add(proj)
    doc = Document(id=str(uuid.uuid4()), project_id=proj.id, filename="plano_piso_10.pdf", page_count=1)
    db_session.add(doc)
    sheet = DocumentSheet(id=str(uuid.uuid4()), document_id=doc.id, sheet_number=1, sheet_code="ARQ-10")
    db_session.add(sheet)
    
    tb = TitleBlockExtraction(
        id=str(uuid.uuid4()),
        sheet_id=sheet.id,
        sheet_code="ARQ-10",
        scale_text="1:50",
        revision="C",
        discipline="architecture"
    )
    db_session.add(tb)

    rule_def = db_session.query(RuleDefinition).first()

    finding1 = RuleFinding(
        id=str(uuid.uuid4()),
        document_id=doc.id,
        sheet_id=sheet.id,
        rule_id=rule_def.id if rule_def else None,
        rule_code="RULE_DOOR_COUNT_MATCH_V1",
        rule_name="Door Count Match",
        category="cross_reconciliation",
        severity="high",
        status="confirmed",
        title="Discrepancia de Conteo de Puertas (Delta: -2)",
        description="Se detectaron 8 puertas pero el cuadro declara 10.",
        expected_value=10,
        observed_value=8,
        delta=-2,
        recommendation="Actualizar rotulación en planta."
    )
    db_session.add(finding1)

    res = FindingResolution(
        id=str(uuid.uuid4()),
        finding_id=finding1.id,
        resolution_type="confirmed",
        resolved_by="auditor_lead",
        notes="Confirmada omisión de 2 puertas de acceso secundario"
    )
    db_session.add(res)

    trace = DecisionTrace(
        id=str(uuid.uuid4()),
        trace_type="rule_evaluation",
        entity_type="RuleExecution",
        entity_id="exec-123",
        confidence=0.95,
        decision_status="review_required",
        explanation="Evaluación de regla con delta de conteo",
        document_id=doc.id,
        sheet_id=sheet.id
    )
    db_session.add(trace)
    db_session.commit()

    svc = ReportingService(db_session)
    report = svc.generate_sheet_report(sheet_id=sheet.id, generated_by="auditor_test")

    assert report is not None
    assert report.status == "completed"
    assert report.report_scope == "sheet"
    assert report.summary["total_findings"] == 1
    assert report.summary["by_severity"]["high"] == 1
    assert report.summary["by_status"]["confirmed"] == 1
    assert report.manifest_hash is not None

    # Verificar que los archivos físicos se hayan creado
    pdf_abs = os.path.join(settings.BASE_DIR, report.artifact_pdf_path)
    json_abs = os.path.join(settings.BASE_DIR, report.artifact_json_path)
    zip_abs = os.path.join(settings.BASE_DIR, report.artifact_bundle_path)

    assert os.path.exists(pdf_abs)
    assert os.path.exists(json_abs)
    assert os.path.exists(zip_abs)

    # Verificar contenido de JSON export
    with open(json_abs, "r", encoding="utf-8") as f:
        json_data = json.load(f)
    assert json_data["metadata"]["report_id"] == report.id
    assert json_data["target"]["sheet_code"] == "ARQ-10"
    assert len(json_data["findings"]) == 1
    assert json_data["findings"][0]["rule_code"] == "RULE_DOOR_COUNT_MATCH_V1"
    assert json_data["findings"][0]["resolutions"][0]["resolved_by"] == "auditor_lead"

    # Verificar contenido del ZIP bundle
    with zipfile.ZipFile(zip_abs, "r") as zf:
        namelist = zf.namelist()
        assert any("audit-report.pdf" in n for n in namelist)
        assert any("audit-report.json" in n for n in namelist)
        assert any("manifest.json" in n for n in namelist)
        assert any("findings.json" in n for n in namelist)

def test_api_report_endpoints_and_download(client, db_session):
    proj = Project(id=str(uuid.uuid4()), name="Centro Cívico", code="CC-01")
    db_session.add(proj)
    doc = Document(id=str(uuid.uuid4()), project_id=proj.id, filename="civico.pdf", page_count=1)
    db_session.add(doc)
    sheet = DocumentSheet(id=str(uuid.uuid4()), document_id=doc.id, sheet_number=1, sheet_code="ARQ-01")
    db_session.add(sheet)
    db_session.commit()

    # 1. Generar reporte síncrono
    resp_create = client.post(f"/api/v1/reports/sheets/{sheet.id}", json={"generated_by": "qa_tester"})
    assert resp_create.status_code == 200
    report_data = resp_create.json()
    report_id = report_data["id"]

    # 2. Consultar listado y detalle
    resp_list = client.get("/api/v1/reports")
    assert resp_list.status_code == 200
    assert len(resp_list.json()) >= 1

    resp_detail = client.get(f"/api/v1/reports/{report_id}")
    assert resp_detail.status_code == 200
    assert resp_detail.json()["id"] == report_id

    # 3. Consultar manifiesto
    resp_manifest = client.get(f"/api/v1/reports/{report_id}/manifest")
    assert resp_manifest.status_code == 200
    manifest = resp_manifest.json()
    assert "files" in manifest["manifest_json"]
    assert len(manifest["manifest_json"]["files"]) >= 4

    # 4. Descargar PDF, JSON y Bundle ZIP
    resp_pdf = client.get(f"/api/v1/reports/{report_id}/download/pdf")
    assert resp_pdf.status_code == 200
    assert resp_pdf.headers["content-type"] == "application/pdf"

    resp_json = client.get(f"/api/v1/reports/{report_id}/download/json")
    assert resp_json.status_code == 200
    assert resp_json.headers["content-type"] == "application/json"

    resp_zip = client.get(f"/api/v1/reports/{report_id}/download/bundle")
    assert resp_zip.status_code == 200
    assert resp_zip.headers["content-type"] == "application/zip"

    # 5. Encolar reporte asíncrono
    resp_async = client.post(f"/api/v1/reports/sheets/{sheet.id}/async")
    assert resp_async.status_code == 202
    assert "job_id" in resp_async.json()


def test_simple_pdf_canvas_encoding_and_wrap_text():
    from app.services.reporting.pdf_renderer import SimplePdfCanvas
    import fitz

    # 1. wrap_text logic
    assert SimplePdfCanvas.wrap_text(None) == []
    assert SimplePdfCanvas.wrap_text("") == []
    sample_text = "Esta es una descripción técnica muy extensa sobre el piping y la viñeta del documento que debe partirse limpiamente sin cortar palabras a la mitad."
    lines = SimplePdfCanvas.wrap_text(sample_text, max_chars_per_line=30)
    assert len(lines) > 1
    for l in lines:
        assert len(l) <= 35  # tolerancia por palabras indivisibles

    # 2. Encoding /WinAnsiEncoding
    canvas = SimplePdfCanvas()
    canvas.new_page()
    spanish_phrase = "Auditoría Técnica: Catálogo de Símbolos en Español -- á é í ó ú ñ Ñ ¿ ¡"
    canvas.draw_text(spanish_phrase, 50, 700, font_size=12, font="Helvetica-Bold")
    pdf_bytes = canvas.build_pdf_bytes()

    assert b"/Encoding /WinAnsiEncoding" in pdf_bytes

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    extracted_text = doc[0].get_text()
    assert "Auditoría Técnica" in extracted_text
    assert "Catálogo de Símbolos" in extracted_text
    assert "á é í ó ú ñ Ñ ¿ ¡" in extracted_text


def test_gen_doc_001_rule_sheet_identification():
    from app.services.rules.piping_rules import GenDoc001Rule
    from app.services.rules.contracts import RuleInput

    rule = GenDoc001Rule()

    class MockDoc:
        filename = "P_ID_001_AREAS.pdf"

    class MockSheet:
        sheet_number = 2
        sheet_code = "SHEET-PID-02"
        title = "Planta de Proceso"

    class MockTitleBlock:
        sheet_code = None  # Falta
        sheet_title = "Planta de Proceso"
        revision = None  # Falta

    inputs = RuleInput(
        document_id="doc-123",
        sheet_id="sheet-456",
        document=MockDoc(),
        sheet=MockSheet(),
        title_block=MockTitleBlock()
    )

    result = rule.evaluate(inputs)
    assert result.status == "failed"
    assert "P_ID_001_AREAS.pdf" in result.title or "P_ID_001_AREAS.pdf" in result.description
    assert "Lámina 2" in result.title or "Lámina 2" in result.description
    assert "sheet_code" in result.title
    assert "revision" in result.title


def test_export_service_pdf_generation_tabular(tmp_path):
    """Verifica que _generate_pdf genere un PDF estructurado con tablas para simbología y resumen compacto de NOT_EVALUABLE."""
    import fitz
    from app.services.review.export_service import ReviewExportService

    pdf_output = str(tmp_path / "test_tabular_audit.pdf")

    run_data = {
        "review_run_id": "run-test-tabular-001",
        "project_code": "PRJ-TEST",
        "project_name": "Planta Minera Concentradora",
        "discipline_code": "PIPING",
        "discipline_name": "Piping & Procesos",
        "topic_code": "TOPIC_PIPING_SPECS",
        "execution_mode": "production",
        "requested_at": "2026-10-10T12:00:00Z",
        "baseline_catalog_version": "V2.0",
        "stats_summary": {
            "passed": 2,
            "failed": 0,
            "warning": 1,
            "not_evaluable": 35
        },
        "steps": [
            {"phase": 1, "phase_name": "Ingesta", "step_type": "intake", "status": "succeeded"},
            {"phase": 2, "phase_name": "Detección Símbolos", "step_type": "symbols", "status": "succeeded"},
            {"phase": 3, "phase_name": "Evaluación Reglas", "step_type": "rules", "status": "succeeded"},
        ],
        "findings": [],
        "symbol_inventory": {
            "metrics": {
                "valid_symbol_occurrences": 8,
                "inventory_groups": 2,
                "figures_excluded": 0,
                "production_coverage": 1.0,
                "sandbox_coverage": 0.0,
                "unknown_rate": 0.0,
                "recognized_production": 8,
                "recognized_sandbox": 0,
                "unknown": 0
            },
            "groups": [
                {
                    "display_code": "PIP-VALVE-GATE",
                    "canonical_name": "Válvula de Compuerta",
                    "catalog_status": "production",
                    "standard_reference": "ASME B16.10 / ASME B31.3",
                    "total_occurrences": 5,
                    "occurrences_by_document": {"PID-001.pdf": 5},
                    "confidence_summary": {"avg": 0.94},
                    "requires_human_review": False,
                    "explanation": "Válvula de compuerta manual bridadada estándar."
                },
                {
                    "display_code": "PIP-VALVE-CHECK",
                    "canonical_name": "Válvula de Retención (Check)",
                    "catalog_status": "production",
                    "standard_reference": "API 594",
                    "total_occurrences": 3,
                    "occurrences_by_document": {"PID-001.pdf": 3},
                    "confidence_summary": {"avg": 0.89},
                    "requires_human_review": False,
                    "explanation": "Válvula check oscilante."
                }
            ],
            "occurrences": [
                {
                    "group_display_code": "PIP-VALVE-GATE",
                    "document_name": "PID-001.pdf",
                    "sheet_name": "Lámina 1",
                    "page_number": 1,
                    "symbol_crop_bbox": [0.12, 0.34, 0.16, 0.38],
                    "detected_tag_or_code": "V-101",
                    "geometric_confidence": 0.96
                }
            ]
        },
        "executions": [
            {
                "rule_code": "RULE_PIP_01",
                "rule_name": "Verificación de Tags de Válvulas",
                "status": "PASSED"
            }
        ] + [
            {
                "rule_code": f"RULE_LEGACY_{i:03d}",
                "rule_name": f"Regla Inactiva {i}",
                "status": "NOT_EVALUABLE",
                "not_evaluable_reason_code": "NO_APPROVED_RULES",
                "recommended_action": "Activar ejecutor en catálogo"
            }
            for i in range(1, 36)
        ]
    }

    ReviewExportService._generate_pdf(run_data, pdf_output)

    assert os.path.exists(pdf_output)
    assert os.path.getsize(pdf_output) > 2000

    doc = fitz.open(pdf_output)
    assert doc.page_count >= 2
    full_text = "".join([p.get_text() for p in doc])

    # Verificar que el PDF incluye la tabla de simbología con sus códigos
    assert "Inventario Consolidado de Simbología" in full_text
    assert "PIP-VALVE-GATE" in full_text
    assert "PIP-VALVE-CHECK" in full_text
    assert "Válvula de Compuerta" in full_text
    # Verificar que las reglas NOT_EVALUABLE fueron resumidas y no causaron 10 páginas
    assert "Resumen de Reglas Sin Ejecutor Activo (35 reglas)" in full_text
    assert "NO_APPROVED_RULES: 35" in full_text
    assert doc.page_count <= 5  # No 15 páginas como antes


def test_export_service_canonical_name_propagation_from_template(db_session, tmp_path):
    """
    Fase 1: Verifica que para un DetectedSymbol con matching_status='matched' y un
    SymbolTemplate asociado con canonical_name poblado, la exportación (XLSX y PDF)
    propaga el canonical_name real a la fila generada en vez de '-'.
    """
    import openpyxl
    import fitz
    from app.db.models.document_memory import DetectedSymbol
    from app.db.models.template_memory import SymbolTemplate
    from app.db.models.decision_memory import ReviewRun, ReviewRunDocument
    from app.services.review.export_service import ReviewExportService

    # 1. Crear proyecto, documento y lámina
    proj = Project(id=str(uuid.uuid4()), name="Proyecto Test Canónico", code="PRJ-CAN-01")
    db_session.add(proj)
    doc = Document(id=str(uuid.uuid4()), project_id=proj.id, filename="PID-CANONICO-001.pdf", page_count=1)
    db_session.add(doc)
    sheet = DocumentSheet(id=str(uuid.uuid4()), document_id=doc.id, sheet_number=1, sheet_code="PID-01")
    db_session.add(sheet)

    # 2. Crear SymbolTemplate con canonical_name explícito
    tmpl = SymbolTemplate(
        id=str(uuid.uuid4()),
        symbol_class="control_valve_diaphragm",
        display_name="• Control Valve with Diaphragm Actuator",
        canonical_name="Válvula de Control Diafragma",
        canonical_code="PIP-VALVE-CTRL-DIA",
        category="valve",
        discipline="piping",
        standard_reference="ISA 5.1 / ASME B16.34",
        status="active"
    )
    db_session.add(tmpl)
    db_session.commit()

    # 3. Crear DetectedSymbol con matching_status="matched" y matched_template_id
    sym = DetectedSymbol(
        id=str(uuid.uuid4()),
        document_id=doc.id,
        sheet_id=sheet.id,
        symbol_type="control_valve_diaphragm",
        discipline="piping",
        matching_status="matched",
        detection_status="detected",
        matched_template_id=tmpl.id,
        matched_library_entry_id=tmpl.id,
        confidence=0.96,
        geometric_confidence=0.95,
        classification="symbol",
        bbox=[100, 100, 150, 150],
        bbox_normalized=[0.1, 0.1, 0.15, 0.15]
    )
    db_session.add(sym)

    # 4. Crear ReviewRun y vincular documento
    run = ReviewRun(
        id=str(uuid.uuid4()),
        project_id=proj.id,
        run_name="Auditoría Test Propagación Canónica",
        execution_mode="production",
        status="succeeded",
        summary_stats={"passed": 1, "failed": 0, "warning": 0, "not_evaluable": 0}
    )
    db_session.add(run)
    db_session.commit()

    r_doc = ReviewRunDocument(
        review_run_id=run.id,
        document_id=doc.id,
        status="included",
        document_role="primary"
    )
    db_session.add(r_doc)
    db_session.commit()

    # 5. Generar reporte XLSX
    report_xlsx = ReviewExportService.create_report(
        db=db_session,
        review_run_id=run.id,
        export_format="xlsx"
    )
    assert report_xlsx is not None
    assert os.path.exists(report_xlsx.artifact_path)

    wb = openpyxl.load_workbook(report_xlsx.artifact_path)
    assert "Resumen de Inventario" in wb.sheetnames
    ws_inv = wb["Resumen de Inventario"]

    # Fila 2 corresponde al primer grupo de símbolos catalogados
    code_val = ws_inv.cell(row=2, column=1).value
    canonical_val = ws_inv.cell(row=2, column=2).value
    std_val = ws_inv.cell(row=2, column=4).value

    assert code_val == "PIP-VALVE-CTRL-DIA"
    assert canonical_val == "Válvula de Control Diafragma"
    assert canonical_val != "-"
    assert "ISA 5.1" in str(std_val)

    # 6. Generar reporte PDF
    report_pdf = ReviewExportService.create_report(
        db=db_session,
        review_run_id=run.id,
        export_format="pdf"
    )
    assert report_pdf is not None
    assert os.path.exists(report_pdf.artifact_path)

    pdf_doc = fitz.open(report_pdf.artifact_path)
    full_pdf_text = "".join([p.get_text() for p in pdf_doc])

    assert "PIP-VALVE-CTRL-DIA" in full_pdf_text
    assert "Válvula de Control Diafragma" in full_pdf_text


