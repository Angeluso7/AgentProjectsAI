import os
import uuid
import json
import pytest
from datetime import datetime
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.db.session import SessionLocal
from app.db.models.core import User, Organization, OrganizationMembership, Project
from app.db.models.decision_memory import (
    ReviewDiscipline,
    ReviewTopic,
    ReviewRun,
    ReviewRunDocument,
    ReviewRunStep,
    SymbolInventoryGroup,
    ReviewReport
)
from app.db.models.document_memory import (
    Document,
    DocumentSheet,
    DetectedSymbol
)
from app.db.models.template_memory import SymbolTemplate
from app.db.models.symbol_catalog import SymbolTemplateVersion
from app.core.security import create_access_token
from app.services.symbols.symbol_inventory_service import SymbolInventoryService
from app.services.review.export_service import ReviewExportService
from app.services.symbols.geometric_validator import (
    compute_symbol_crop_bbox,
    compute_occurrence_context_crop_bbox
)

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


@pytest.fixture
def inventory_test_env():
    db: Session = SessionLocal()
    suffix = str(uuid.uuid4())[:8]

    org = Organization(id=f"org-inv-{suffix}", name=f"Org Inv {suffix}", slug=f"org-inv-{suffix}")
    user = User(
        id=f"user-inv-{suffix}",
        email=f"auditor-inv-{suffix}@test.com",
        display_name=f"Auditor Inv {suffix}",
        password_hash="fake",
        is_active=True
    )
    db.add_all([org, user])
    db.commit()

    mem = OrganizationMembership(
        organization_id=org.id,
        user_id=user.id,
        role="admin"
    )
    db.add(mem)

    project = Project(
        id=f"prj-inv-{suffix}",
        organization_id=org.id,
        code=f"PRJ-INV-{suffix}",
        name=f"Proyecto Inv {suffix}",
        status="active"
    )
    db.add(project)
    db.commit()

    disc = ReviewDiscipline(
        id=f"disc-piping-{suffix}",
        code="PIPING",
        name="Cañerías / Piping",
        order_index=1,
        is_active=True
    )
    topic = ReviewTopic(
        id=f"topic-pid-{suffix}",
        discipline_id=disc.id,
        code="PID_SYMBOLS",
        name="Simbología P&ID",
        order_index=1,
        is_active=True
    )
    db.add_all([disc, topic])
    db.commit()

    # Documentos
    doc_legend = Document(
        id=f"doc-leg-{suffix}",
        organization_id=org.id,
        project_id=project.id,
        filename="000-LEGEND.pdf",
        file_path=f"/storage/docs/000-LEGEND-{suffix}.pdf",
        file_hash_sha256=f"hash-leg-{suffix}",
        file_size_bytes=102400,
        page_count=1,
        status="ready"
    )
    doc_pid = Document(
        id=f"doc-pid-{suffix}",
        organization_id=org.id,
        project_id=project.id,
        filename="001-PID-AREA-100.pdf",
        file_path=f"/storage/docs/001-PID-{suffix}.pdf",
        file_hash_sha256=f"hash-pid-{suffix}",
        file_size_bytes=204800,
        page_count=1,
        status="ready"
    )
    db.add_all([doc_legend, doc_pid])
    db.commit()

    sheet_leg = DocumentSheet(
        id=f"sheet-leg-{suffix}",
        document_id=doc_legend.id,
        sheet_number=1,
        title="Leyenda Simbología",
        width_px=2400,
        height_px=1800,
        dpi=150
    )
    sheet_pid = DocumentSheet(
        id=f"sheet-pid-{suffix}",
        document_id=doc_pid.id,
        sheet_number=1,
        title="P&ID Hoja 1",
        width_px=2400,
        height_px=1800,
        dpi=150
    )
    db.add_all([sheet_leg, sheet_pid])
    db.commit()

    # Plantilla de Catálogo (para símbolo reconocido)
    tmpl = SymbolTemplate(
        id=f"tmpl-valve-{suffix}",
        organization_id=org.id,
        symbol_class="gate_valve",
        display_name="Válvula Compuerta Manual",
        canonical_code="SYM-GATE_VALVE",
        canonical_name="Válvula Compuerta Manual",
        technical_function="Aislamiento on/off de línea de proceso",
        standard_reference="ISA-5.1 / PIP PNC00001",
        status="active"
    )
    db.add(tmpl)
    db.commit()

    tmpl_ver = SymbolTemplateVersion(
        id=f"ver-valve-{suffix}",
        symbol_template_id=tmpl.id,
        version_number=1,
        approval_status="approved"
    )
    db.add(tmpl_ver)
    db.commit()

    # Corrida ReviewRun
    run = ReviewRun(
        id=f"run-inv-{suffix}",
        organization_id=org.id,
        project_id=project.id,
        discipline_id=disc.id,
        topic_id=topic.id,
        execution_mode="production",
        run_name=f"Auditoría Simbología {suffix}",
        status="succeeded",
        requested_by=user.email,
        summary_stats={"passed": 5, "failed": 1, "warning": 0, "not_evaluable": 0}
    )
    db.add(run)
    db.commit()

    rd1 = ReviewRunDocument(
        id=f"rd1-{suffix}",
        review_run_id=run.id,
        document_id=doc_legend.id,
        inclusion_reason="legend",
        document_role="legend",
        status="included"
    )
    rd2 = ReviewRunDocument(
        id=f"rd2-{suffix}",
        review_run_id=run.id,
        document_id=doc_pid.id,
        inclusion_reason="pid",
        document_role="pid",
        status="included"
    )
    db.add_all([rd1, rd2])
    db.commit()

    # Símbolos detectados
    # 1. Reconocido productivo (2 ocurrencias: 1 en leyenda, 1 en PID)
    s1 = DetectedSymbol(
        id=f"sym1-{suffix}",
        document_id=doc_legend.id,
        sheet_id=sheet_leg.id,
        symbol_type="gate_valve",
        classification="symbol",
        matched_template_id=tmpl.id,
        matched_template_version_id=tmpl_ver.id,
        matching_status="matched",
        geometric_confidence=0.96,
        bbox=[240, 360, 360, 450],
        bbox_normalized=[0.10, 0.20, 0.15, 0.25],
        inner_drawing_bbox=[0.105, 0.205, 0.145, 0.245],
        detected_tag_or_code="V-101"
    )
    s2 = DetectedSymbol(
        id=f"sym2-{suffix}",
        document_id=doc_pid.id,
        sheet_id=sheet_pid.id,
        symbol_type="gate_valve",
        classification="symbol",
        matched_template_id=tmpl.id,
        matched_template_version_id=tmpl_ver.id,
        matching_status="matched",
        geometric_confidence=0.94,
        bbox=[720, 720, 840, 810],
        bbox_normalized=[0.30, 0.40, 0.35, 0.45],
        inner_drawing_bbox=[0.305, 0.405, 0.345, 0.445],
        detected_tag_or_code="V-102"
    )

    # 2. Desconocido (unknown_symbol) con geometría vectorial válida
    s3 = DetectedSymbol(
        id=f"sym3-{suffix}",
        document_id=doc_pid.id,
        sheet_id=sheet_pid.id,
        symbol_type="custom_regulator",
        classification="symbol",
        matched_template_id=None,
        matched_template_version_id=None,
        matching_status="unmatched",
        geometric_confidence=0.91,
        bbox=[1200, 1080, 1320, 1170],
        bbox_normalized=[0.50, 0.60, 0.55, 0.65],
        inner_drawing_bbox=[0.505, 0.605, 0.545, 0.645],
        detected_tag_or_code="PCV-201"
    )
    s4 = DetectedSymbol(
        id=f"sym4-{suffix}",
        document_id=doc_pid.id,
        sheet_id=sheet_pid.id,
        symbol_type="custom_regulator",
        classification="symbol",
        matched_template_id=None,
        matched_template_version_id=None,
        matching_status="unmatched",
        geometric_confidence=0.89,
        bbox=[1680, 1260, 1800, 1350],
        bbox_normalized=[0.70, 0.70, 0.75, 0.75],
        inner_drawing_bbox=[0.705, 0.705, 0.745, 0.745],
        detected_tag_or_code="PCV-202"
    )

    # 3. Figura excluida
    s5 = DetectedSymbol(
        id=f"sym5-{suffix}",
        document_id=doc_legend.id,
        sheet_id=sheet_leg.id,
        symbol_type="flow_diagram_box",
        classification="figure",
        matching_status="excluded",
        geometric_confidence=0.99,
        bbox=[1920, 180, 2280, 450],
        bbox_normalized=[0.80, 0.10, 0.95, 0.25]
    )

    db.add_all([s1, s2, s3, s4, s5])
    db.commit()

    token = create_access_token(subject=user.id, email=user.email, extra_claims={"organization_id": org.id})

    yield {
        "db": db,
        "org": org,
        "user": user,
        "project": project,
        "run": run,
        "tmpl": tmpl,
        "tmpl_ver": tmpl_ver,
        "symbols": [s1, s2, s3, s4, s5],
        "token": token
    }

    db.close()


def test_symbol_inventory_double_crops_invariants(inventory_test_env):
    """
    Invariante Fundamental:
    A. symbol_crop: inner_drawing_bbox + 3 mm
    B. occurrence_context_crop: symbol_crop_bbox + 15 mm, limitado a [0, 0, 1, 1]
    """
    sym = inventory_test_env["symbols"][0]
    pw, ph = 841.0, 595.0

    SymbolInventoryService.ensure_crops_for_symbol(sym, None, pw, ph)

    assert sym.symbol_crop_bbox is not None
    assert len(sym.symbol_crop_bbox) == 4
    assert sym.occurrence_context_crop_bbox is not None
    assert len(sym.occurrence_context_crop_bbox) == 4

    # El crop de contexto (15 mm) debe envolver estrictamente al crop de símbolo (3 mm)
    sc = sym.symbol_crop_bbox
    cc = sym.occurrence_context_crop_bbox

    assert cc[0] <= sc[0]
    assert cc[1] <= sc[1]
    assert cc[2] >= sc[2]
    assert cc[3] >= sc[3]

    # Límites estrictos dentro de la página [0, 0, 1, 1]
    assert cc[0] >= 0.0 and cc[1] >= 0.0
    assert cc[2] <= 1.0 and cc[3] <= 1.0


def test_symbol_inventory_grouping_and_unknowns(inventory_test_env):
    """
    Valida:
    1. Agrupación por template para símbolos catalogados (recognized_production).
    2. Agrupación por firma para símbolos desconocidos (unknown_symbol con U-001).
    3. Exclusión de figuras fuera del conteo canónico (figure_excluded).
    4. Conteo por documento y hoja.
    """
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]

    groups, metrics = SymbolInventoryService.build_run_inventory(db, run, force_rebuild=True)

    assert len(groups) >= 3  # gate_valve, custom_regulator (U-001), figure_excluded

    prod_group = next(g for g in groups if g.catalog_status == "recognized_production")
    assert prod_group.display_code == "SYM-GATE_VALVE"
    assert prod_group.total_occurrences == 2
    assert "000-LEGEND.pdf" in prod_group.occurrences_by_document
    assert "001-PID-AREA-100.pdf" in prod_group.occurrences_by_document

    unk_group = next(g for g in groups if g.catalog_status == "unknown_symbol")
    assert unk_group.display_code.startswith("U-")
    assert unk_group.total_occurrences == 2
    assert unk_group.requires_human_review is True

    fig_group = next(g for g in groups if g.catalog_status == "figure_excluded")
    assert fig_group.display_code.startswith("FIG-")
    assert fig_group.total_occurrences == 1

    # Validar métricas consolidadas
    assert metrics["valid_symbol_occurrences"] == 4  # 2 gate valves + 2 custom regulators
    assert metrics["figures_excluded"] == 1
    assert metrics["recognized_production"] == 2
    assert metrics["unknown"] == 2
    assert metrics["production_coverage"] == 0.5  # 2/4 = 50%
    assert metrics["unknown_rate"] == 0.5         # 2/4 = 50%


def test_grouped_unknown_findings_rule_of_business(inventory_test_env):
    """
    Regla de Negocio 6: NO crear un hallazgo individual por cada ocurrencia desconocida.
    Debe crearse UN solo hallazgo agrupado por familia desconocida U-001.
    """
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]

    groups, _ = SymbolInventoryService.build_run_inventory(db, run, force_rebuild=False)
    unknown_groups = [g for g in groups if g.catalog_status == "unknown_symbol"]

    findings = SymbolInventoryService.create_grouped_unknown_findings(db, run, unknown_groups)

    assert len(findings) == 1  # Exactamente 1 hallazgo agrupado para U-001
    finding = findings[0]
    assert "U-001" in finding.title or "U-" in finding.title
    assert finding.evidence_refs["total_occurrences"] == 2
    assert len(finding.evidence_refs["occurrence_ids"]) == 2


def test_symbol_inventory_api_endpoints(inventory_test_env):
    """
    Verifica los endpoints:
    1. GET /api/v1/review/runs/{run_id}/symbol-inventory inicial (corrida sin inventario generado previo -> unavailable)
    2. POST /api/v1/review/runs/{run_id}/symbol-inventory/regenerate (genera versión v1)
    3. Idempotencia de regeneración
    4. GET /api/v1/review/runs/{run_id}/symbol-inventory (Tabla 1 disponible con grupos)
    5. GET /api/v1/review/runs/{run_id}/symbol-inventory/groups/{group_id}/occurrences (Tabla 2)
    """
    run = inventory_test_env["run"]
    token = inventory_test_env["token"]
    client = TestClient(app)

    # 1. Tabla 1 antes de regenerar: estado unavailable seguro
    res_initial = client.get(
        f"/api/v1/review/runs/{run.id}/symbol-inventory",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res_initial.status_code == 200, res_initial.text
    data_initial = res_initial.json()
    assert data_initial["status"] == "unavailable"
    assert data_initial["can_generate"] is True
    assert data_initial["metrics"]["valid_symbol_occurrences"] == 0
    assert len(data_initial["groups"]) == 0

    # 2. POST regenerate: genera el inventario
    res_regen = client.post(
        f"/api/v1/review/runs/{run.id}/symbol-inventory/regenerate",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res_regen.status_code == 200, res_regen.text
    data_regen = res_regen.json()
    assert data_regen["status"] == "available"
    assert data_regen["inventory_version"] == "v1"
    assert data_regen["metrics"]["valid_symbol_occurrences"] == 4
    assert len(data_regen["groups"]) >= 2
    assert data_regen["inventory_source_snapshot_hash"] is not None

    # 3. Idempotencia: segunda llamada a regenerar con el mismo snapshot
    res_regen_idem = client.post(
        f"/api/v1/review/runs/{run.id}/symbol-inventory/regenerate",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res_regen_idem.status_code == 200, res_regen_idem.text
    data_regen_idem = res_regen_idem.json()
    assert data_regen_idem["inventory_version"] == "v1"
    assert data_regen_idem["inventory_source_snapshot_hash"] == data_regen["inventory_source_snapshot_hash"]

    # 4. Tabla 1: Inventario Consolidado tras generación
    res_inv = client.get(
        f"/api/v1/review/runs/{run.id}/symbol-inventory",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res_inv.status_code == 200, res_inv.text
    data_inv = res_inv.json()

    assert data_inv["status"] == "available"
    assert "metrics" in data_inv
    assert "groups" in data_inv
    assert "excluded_groups" in data_inv
    assert data_inv["metrics"]["valid_symbol_occurrences"] == 4
    assert len(data_inv["groups"]) >= 2

    target_group = data_inv["groups"][0]
    group_id = target_group["id"]

    # 5. Tabla 2: Ocurrencias con Doble Crop
    res_occ = client.get(
        f"/api/v1/review/runs/{run.id}/symbol-inventory/groups/{group_id}/occurrences",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res_occ.status_code == 200, res_occ.text
    data_occ = res_occ.json()

    assert len(data_occ) == target_group["total_occurrences"]
    for occ in data_occ:
        assert occ["symbol_crop_bbox"] is not None
        assert occ["occurrence_context_crop_bbox"] is not None
        assert occ["context_margin_mm"] == 15.0
        assert occ["navigation_context"]["document_id"] is not None


def test_symbol_inventory_regenerate_blocked_when_running(inventory_test_env):
    """Verifica que regenerar inventario falle con HTTP 400 si la corrida está en progreso."""
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]
    token = inventory_test_env["token"]
    client = TestClient(app)

    run.status = "running"
    db.commit()

    res = client.post(
        f"/api/v1/review/runs/{run.id}/symbol-inventory/regenerate",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res.status_code == 400
    assert "ejecución" in res.json()["detail"].lower()

    # Restaurar status
    run.status = "succeeded"
    db.commit()


def test_symbol_inventory_exports_pdf_and_xlsx(inventory_test_env):
    """
    Valida la incorporación de Tabla 1 y Tabla 2 en los reportes persistidos XLSX y PDF.
    """
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]

    # 1. Export XLSX
    rep_xlsx = ReviewExportService.create_report(db, run.id, export_format="xlsx")
    assert rep_xlsx.status == "ready"
    assert os.path.exists(rep_xlsx.artifact_path)

    if OPENPYXL_AVAILABLE:
        wb = openpyxl.load_workbook(rep_xlsx.artifact_path)
        sheet_names = wb.sheetnames
        assert "Resumen de Inventario" in sheet_names
        assert "Ocurrencias y Localizaciones" in sheet_names

        ws_inv = wb["Resumen de Inventario"]
        assert ws_inv.max_row >= 3  # Headers + al menos 2 grupos

        ws_occ = wb["Ocurrencias y Localizaciones"]
        assert ws_occ.max_row >= 5  # Headers + 4 ocurrencias

    # 2. Export PDF
    rep_pdf = ReviewExportService.create_report(db, run.id, export_format="pdf")
    assert rep_pdf.status == "ready"
    assert os.path.exists(rep_pdf.artifact_path)
    assert len(rep_pdf.sha256) == 64
    with open(rep_pdf.artifact_path, "rb") as f:
        content = f.read()
        assert content.startswith(b"%PDF")


def test_symbol_inventory_api_never_returns_null_contract(inventory_test_env):
    """
    CONTRATO OBLIGATORIO:
    GET /review/runs/{run_id} y GET /review/runs/{run_id}/symbol-inventory
    NUNCA deben retornar symbol_inventory = null bajo ningún estado:
    - corrida histórica (unavailable)
    - corrida en ejecución (pending)
    - corrida completada con o sin grupos (available)
    """
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]
    token = inventory_test_env["token"]
    client = TestClient(app)

    # 1. Estado Histórico / Inicial (unavailable)
    # A. Detalle de corrida
    res_run = client.get(f"/api/v1/review/runs/{run.id}", headers={"Authorization": f"Bearer {token}"})
    assert res_run.status_code == 200
    run_data = res_run.json()
    assert run_data["symbol_inventory"] is not None
    inv = run_data["symbol_inventory"]
    assert inv["status"] in ["available", "pending", "unavailable", "failed"]
    assert isinstance(inv["metrics"], dict)
    assert isinstance(inv["groups"], list)
    assert isinstance(inv["excluded_groups"], list)
    assert "valid_symbol_occurrences" in inv["metrics"]

    # B. Endpoint dedicado
    res_inv = client.get(f"/api/v1/review/runs/{run.id}/symbol-inventory", headers={"Authorization": f"Bearer {token}"})
    assert res_inv.status_code == 200
    inv_dedicated = res_inv.json()
    assert inv_dedicated is not None
    assert inv_dedicated["status"] in ["available", "pending", "unavailable", "failed"]
    assert isinstance(inv_dedicated["groups"], list)
    assert isinstance(inv_dedicated["metrics"], dict)

    # 2. Estado en ejecución (pending)
    run.status = "running"
    db.commit()

    res_running = client.get(f"/api/v1/review/runs/{run.id}", headers={"Authorization": f"Bearer {token}"})
    assert res_running.status_code == 200
    inv_running = res_running.json()["symbol_inventory"]
    assert inv_running is not None
    assert inv_running["status"] == "pending"
    assert inv_running["reason_code"] == "RUN_IN_PROGRESS"
    assert inv_running["can_generate"] is False

    res_running_inv = client.get(f"/api/v1/review/runs/{run.id}/symbol-inventory", headers={"Authorization": f"Bearer {token}"})
    assert res_running_inv.status_code == 200
    assert res_running_inv.json()["status"] == "pending"

    # Restaurar estado
    run.status = "succeeded"
    db.commit()


def test_symbol_inventory_available_status_with_zero_groups(inventory_test_env):
    """
    SEMÁNTICA DE ESTADO:
    available con 0 grupos:
    La corrida procesó documentos pero no detectó geometrías de símbolos.
    El build finalizó exitosamente (inventory_build_completed: True).
    El estado API debe ser 'available' (no empty como status separado),
    con grupos=[] y valid_symbol_occurrences=0.
    """
    db = inventory_test_env["db"]
    org = inventory_test_env["org"]
    project = inventory_test_env["project"]
    user = inventory_test_env["user"]
    token = inventory_test_env["token"]
    client = TestClient(app)

    empty_run = ReviewRun(
        id=f"run-empty-{uuid.uuid4().hex[:8]}",
        organization_id=org.id,
        project_id=project.id,
        execution_mode="production",
        run_name="Corrida Sin Símbolos",
        status="succeeded",
        requested_by=user.email,
        summary={
            "symbol_inventory_meta": {
                "status": "available",
                "inventory_build_completed": True,
                "inventory_version": "v1",
                "inventory_generated_at": datetime.utcnow().isoformat(),
                "inventory_generated_by": user.email,
                "inventory_source_snapshot_hash": "hash_empty_test",
                "inventory_recalculation_count": 0,
                "inventory_algorithm_version": "v1.0"
            }
        }
    )
    db.add(empty_run)
    db.commit()

    res = client.get(
        f"/api/v1/review/runs/{empty_run.id}/symbol-inventory",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "available"
    assert data["inventory_generated_at"] is not None
    assert len(data["groups"]) == 0
    assert data["metrics"]["valid_symbol_occurrences"] == 0
    assert data["can_generate"] is True


def test_symbol_inventory_tenant_isolation(inventory_test_env):
    """
    AISLAMIENTO MULTI-TENANT:
    Un usuario de una organización distinta NO puede acceder a:
    - GET /review/runs/{run_id} -> 403 Forbidden
    - GET /review/runs/{run_id}/symbol-inventory -> 403 Forbidden
    - POST /review/runs/{run_id}/symbol-inventory/regenerate -> 403 Forbidden
    """
    db = inventory_test_env["db"]
    run1 = inventory_test_env["run"]
    client = TestClient(app)

    # Crear Organización 2 y Usuario 2
    suffix2 = uuid.uuid4().hex[:6]
    org2 = Organization(
        id=f"org2-{suffix2}",
        name=f"Organización Externa {suffix2}",
        slug=f"org2-{suffix2}"
    )
    user2 = User(
        id=f"usr2-{suffix2}",
        email=f"externo-{suffix2}@test.com",
        display_name=f"Usuario Externo {suffix2}",
        password_hash="fake",
        is_active=True
    )
    db.add_all([org2, user2])
    db.commit()

    token2 = create_access_token(subject=user2.id, email=user2.email, extra_claims={"organization_id": org2.id})

    # Intentar acceder a run de org1 con token de org2
    res_details = client.get(
        f"/api/v1/review/runs/{run1.id}",
        headers={"Authorization": f"Bearer {token2}"}
    )
    assert res_details.status_code == 403

    res_inv = client.get(
        f"/api/v1/review/runs/{run1.id}/symbol-inventory",
        headers={"Authorization": f"Bearer {token2}"}
    )
    assert res_inv.status_code == 403

    res_regen = client.post(
        f"/api/v1/review/runs/{run1.id}/symbol-inventory/regenerate",
        headers={"Authorization": f"Bearer {token2}"}
    )
    assert res_regen.status_code == 403


def test_symbol_inventory_version_history_and_report_version_preservation(inventory_test_env):
    """
    REGENERACIÓN VERSIONADA Y PRESERVACIÓN EN INFORMES:
    - Se genera inventario v1.
    - Se crea un informe (PDF/XLSX/JSON) que conserva inventory_version='v1'.
    - Se fuerza una regeneración posterior a versión v2.
    - Se verifica que:
      1. El ReviewRun avanza a v2 con recalculation_count=1 y registro en audit_history.
      2. El informe previamente generado conserva intacta la versión 'v1' utilizada.
    """
    db = inventory_test_env["db"]
    run = inventory_test_env["run"]

    # 1. Generar inventario inicial (v1)
    SymbolInventoryService.build_run_inventory(db, run, force_rebuild=True, requested_by="auditor_v1")
    meta_v1 = (run.summary or {}).get("symbol_inventory_meta", {})
    assert meta_v1["inventory_version"] == "v1"
    assert meta_v1["inventory_recalculation_count"] == 0

    # 2. Generar informe exportable con versión v1
    report_v1 = ReviewExportService.create_report(db, run.id, export_format="json")
    assert report_v1.stats_summary.get("inventory_version") == "v1"

    # Verificar contenido del archivo JSON
    with open(report_v1.artifact_path, "r", encoding="utf-8") as f:
        file_data = json.load(f)
    assert file_data["symbol_inventory"]["version"] == "v1"

    # 3. Forzar regeneración posterior
    SymbolInventoryService.build_run_inventory(db, run, force_rebuild=True, requested_by="auditor_v2")
    db.refresh(run)
    meta_v2 = (run.summary or {}).get("symbol_inventory_meta", {})

    assert meta_v2["inventory_version"] == "v2"
    assert meta_v2["inventory_recalculation_count"] == 1
    assert len(meta_v2["audit_history"]) >= 2
    assert meta_v2["audit_history"][-1]["action"] == "recalculated"
    assert meta_v2["audit_history"][-1]["generated_by"] == "auditor_v2"

    # 4. Verificar que el informe histórico NO fue sobreescrito silenciosamente
    db.refresh(report_v1)
    assert report_v1.stats_summary.get("inventory_version") == "v1"
    with open(report_v1.artifact_path, "r", encoding="utf-8") as f:
        file_data_after = json.load(f)
    assert file_data_after["symbol_inventory"]["version"] == "v1"

